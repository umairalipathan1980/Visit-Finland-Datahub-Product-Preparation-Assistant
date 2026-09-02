"""Coarse, typed tools for the optimized continuous agent workflow.

Claude keeps semantic control.  This module owns paths, subprocess argv,
network boundaries, artifact writes, validation, and compact evidence
packaging so none of those mechanics require general shell or file tools.
"""
from __future__ import annotations

import asyncio
import copy
import functools
import hashlib
import json
import os
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Annotated, Any, Literal
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from claude_agent_sdk import create_sdk_mcp_server, tool
import jsonschema
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from app.telemetry import RunTelemetry
from app.reporting import with_deterministic_appendix

SERVER_NAME = "visit_finland"
TOOL_PREFIX = f"mcp__{SERVER_NAME}__"
MAX_SELECTED_PER_CALL = 20
MAX_RECOVERY_SELECTED = 5
MAX_SELECTED_PER_RUN = 25
MAX_EXPANSION_CALLS = 2
MAX_TOTAL_PAGES = 25
MAX_CONTEXT_PAGE_CHARS = 48_000
MAX_MCP_RESPONSE_CHARS = 52_000
MAX_EVIDENCE_CHUNK_CHARS = 44_000
MAX_TOTAL_CONTEXT_CHARS = 2_400_000
MAX_CANDIDATE_URL_CHARS = 2_000
MAX_CANDIDATE_TEXT_CHARS = 1_000
CANDIDATE_KEYWORDS = (
    "accommodation", "hotel", "room", "contact", "location", "booking", "terms",
    "accessibility", "sustainability", "majoitus", "hotelli", "huone", "yhteys",
    "sijainti", "varaus", "ehdot", "esteettomyys", "vastuullisuus",
)

TOOL_NAMES = [
    "load_skill",
    "prepare_sources",
    "read_context_page",
    "fetch_selected_pages",
    "record_scope",
    "submit_extraction",
    "finalize_outputs",
]
ALLOWED_TOOL_NAMES = [f"{TOOL_PREFIX}{name}" for name in TOOL_NAMES]


class _StrictToolInput(BaseModel):
    model_config = ConfigDict(extra="forbid")


class EmptyToolInput(_StrictToolInput):
    pass


class FetchSelectedPagesInput(_StrictToolInput):
    link_ids: list[Annotated[str, Field(pattern=r"^l[0-9]{3,}$")]] = Field(
        min_length=1, max_length=MAX_SELECTED_PER_CALL,
    )
    reason: str = Field(min_length=1, max_length=500)


class ReadContextPageInput(_StrictToolInput):
    cursor: Annotated[str, Field(pattern=r"^ctx[0-9]{4,}$")]


class ScopeProduct(_StrictToolInput):
    pages: list[Annotated[str, Field(pattern=r"^[pd][0-9]{3,}$")]] = Field(min_length=1)
    summary: str = Field(min_length=1, max_length=2000)


class ResolvedScopeDecision(_StrictToolInput):
    status: Literal["resolved"]
    error_code: None = None
    product: ScopeProduct
    excluded: list[Any]
    additional_products: list[Any]
    out_of_type_facilities: list[Any]


class AmbiguousScopeDecision(_StrictToolInput):
    status: Literal["scope_ambiguous"]
    error_code: Literal["multiple_accommodation_products", "product_boundary_unclear"]
    reason: str = Field(min_length=1, max_length=2000)
    product: ScopeProduct | None = None
    excluded: list[Any]
    additional_products: list[Any]
    out_of_type_facilities: list[Any]


class ScopeInput(_StrictToolInput):
    scope_decision: Annotated[
        ResolvedScopeDecision | AmbiguousScopeDecision,
        Field(discriminator="status"),
    ]


class ExtractionInput(_StrictToolInput):
    extraction: dict[str, Any]


class FinalizeInput(_StrictToolInput):
    review_report: str = Field(min_length=20)


@dataclass
class OptimizedToolState:
    workspace: Path
    package_dir: Path
    telemetry: RunTelemetry
    skill_loaded: bool = False
    prepared: bool = False
    scope_status: str | None = None
    expansion_calls: int = 0
    selected_link_ids: set[str] = field(default_factory=set)
    extraction_submissions: int = 0
    finalized: bool = False
    forced_status: str | None = None
    forced_error_code: str | None = None
    stage_reached: str = "stage_0_bootstrap"
    idempotent_results: dict[str, dict[str, Any]] = field(default_factory=dict)
    active_tools: set[str] = field(default_factory=set)
    artifact_hashes: dict[str, str] = field(default_factory=dict)
    evidence_chars_returned: int = 0
    context_pages: dict[str, dict[str, Any]] = field(default_factory=dict)
    pending_context_cursors: list[str] = field(default_factory=list)
    delivered_context_cursors: set[str] = field(default_factory=set)
    next_context_page_index: int = 1
    queued_candidate_ids: set[str] = field(default_factory=set)
    context_chars_queued: int = 0


def _mark_scope_reasoning_started(state: OptimizedToolState) -> None:
    state.stage_reached = "stage_4"
    state.telemetry.emit({
        "type": "stage",
        "stage": "stage_4",
        "operation": "determine_scope",
        "state": "started",
    })


def _json_result(payload: dict[str, Any], *, is_error: bool = False) -> dict[str, Any]:
    return {
        "content": [{"type": "text", "text": json.dumps(payload, ensure_ascii=False)}],
        "is_error": is_error,
    }


def _validated(model: type[BaseModel], args: dict[str, Any]) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    try:
        return model.model_validate(args).model_dump(), None
    except ValidationError as exc:
        return None, _json_result({
            "status": "error",
            "error_code": "invalid_tool_input",
            "details": exc.errors(include_url=False),
        }, is_error=True)


def _idempotent_handler(state: OptimizedToolState, tool_name: str):
    def decorator(function):
        @functools.wraps(function)
        async def wrapped(args: dict[str, Any]):
            key = f"{tool_name}:{json.dumps(args, sort_keys=True, ensure_ascii=False, separators=(',', ':'))}"
            if key in state.idempotent_results:
                return copy.deepcopy(state.idempotent_results[key])
            if tool_name in state.active_tools:
                return _json_result({
                    "status": "error",
                    "error_code": "operation_in_progress",
                    "tool": tool_name,
                }, is_error=True)
            state.active_tools.add(tool_name)
            try:
                response = await function(args)
            finally:
                state.active_tools.discard(tool_name)
            if not response.get("is_error", False):
                state.idempotent_results[key] = copy.deepcopy(response)
            return response
        return wrapped
    return decorator


def _metric_fields(result: dict[str, Any] | None) -> dict[str, Any]:
    return {key: value for key, value in (result or {}).items() if key != "status"}


def _read_json(path: Path, default: Any = None) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return default


def _write_json_atomic(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = path.with_name(f".{path.name}.tmp")
    temp_path.write_text(json.dumps(value, indent=2, ensure_ascii=False), encoding="utf-8")
    os.replace(temp_path, path)


async def _run_script(state: OptimizedToolState, script_name: str, *args: str) -> dict[str, Any]:
    script = state.package_dir / "scripts" / script_name
    process = await asyncio.create_subprocess_exec(
        sys.executable,
        str(script),
        *args,
        cwd=str(state.package_dir),
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.PIPE,
    )
    stdout_b, stderr_b = await process.communicate()
    stdout = stdout_b.decode("utf-8", errors="replace").strip()
    stderr = stderr_b.decode("utf-8", errors="replace").strip()
    parsed_stdout = None
    if stdout:
        try:
            parsed_stdout = json.loads(stdout.splitlines()[-1])
        except json.JSONDecodeError:
            parsed_stdout = None
    parsed_stderr = None
    if stderr:
        try:
            parsed_stderr = json.loads(stderr.splitlines()[-1])
        except json.JSONDecodeError:
            parsed_stderr = None
    return {
        "returncode": process.returncode,
        "result": parsed_stdout,
        "error": parsed_stderr or ({"message": stderr[-1000:]} if stderr else None),
    }


def _page_records(workspace: Path, source_ids: set[str] | None = None) -> list[dict[str, Any]]:
    index = _read_json(workspace / "work" / "fetched-pages.json", {"pages": []})
    records: list[dict[str, Any]] = []
    for page in index.get("pages", []):
        source_id = page.get("id")
        if not source_id or (source_ids is not None and source_id not in source_ids):
            continue
        meta = _read_json(workspace / "work" / "pages" / f"{source_id}.meta.json", {})
        try:
            text = (workspace / "work" / "pages" / f"{source_id}.txt").read_text(encoding="utf-8")
        except OSError:
            continue
        records.append({
            "source_id": source_id,
            "source_type": "website",
            "url": page.get("url"),
            "final_url": page.get("final_url"),
            "canonical_url": page.get("canonical_url"),
            "content_hash": page.get("content_hash"),
            "title": meta.get("title", page.get("title")),
            "language_hint": meta.get("language_hint", page.get("language_hint")),
            "retrieved_at": meta.get("retrieved_at", page.get("retrieved_at")),
            "content_bytes": meta.get("content_bytes", page.get("content_bytes")),
            "locator": "stored page text",
            "text": text,
            "internal_links": meta.get("internal_link_details", meta.get("internal_links", [])),
            "external_links": meta.get("external_link_details", meta.get("external_links", [])),
        })
    return records


def _document_records(workspace: Path) -> list[dict[str, Any]]:
    parsed = _read_json(workspace / "work" / "parsed-documents.json", {"documents": []})
    records: list[dict[str, Any]] = []
    for doc in parsed.get("documents", []):
        if doc.get("status") != "parsed" or doc.get("ocr_required"):
            continue
        records.append({
            "source_id": doc.get("id"),
            "source_type": "document",
            "file_name": doc.get("file_name"),
            "parsing_method": doc.get("parsing_method"),
            "segments": doc.get("segments", []),
            "text": doc.get("flat_text", ""),
        })
    return records


TRACKING_QUERY_KEYS = {
    "fbclid", "gclid", "gbraid", "wbraid", "gad_source", "mc_cid", "mc_eid",
}


def _display_url(value: Any) -> str | None:
    if not isinstance(value, str) or not value:
        return None
    try:
        parts = urlsplit(value)
        query = [
            (key, item)
            for key, item in parse_qsl(parts.query, keep_blank_values=True)
            if not key.lower().startswith("utm_") and key.lower() not in TRACKING_QUERY_KEYS
        ]
        cleaned = urlunsplit((parts.scheme, parts.netloc, parts.path, urlencode(query, doseq=True), ""))
    except ValueError:
        cleaned = value
    if len(cleaned) <= MAX_CANDIDATE_URL_CHARS:
        return cleaned
    return cleaned[: MAX_CANDIDATE_URL_CHARS - 3] + "..."


def _compact_candidate(candidate: dict[str, Any]) -> dict[str, Any]:
    compact = {
        "id": candidate.get("id"),
        "url": _display_url(candidate.get("url")),
        "anchor_text": str(candidate.get("anchor_text") or "")[:MAX_CANDIDATE_TEXT_CHARS],
        "discovered_on": candidate.get("discovered_on"),
    }
    surrounding = str(candidate.get("surrounding_text") or "")
    if surrounding:
        compact["surrounding_text"] = surrounding[:MAX_CANDIDATE_TEXT_CHARS]
    return compact


def _compact_metadata(value: Any, max_chars: int) -> str | None:
    if value is None:
        return None
    return str(value)[:max_chars]


def _split_text_losslessly(text: str, max_chars: int = MAX_EVIDENCE_CHUNK_CHARS):
    start = 0
    while start < len(text):
        hard_end = min(len(text), start + max_chars)
        end = hard_end
        if hard_end < len(text):
            minimum = start + max_chars // 2
            for separator in ("\n\n", "\n", " "):
                boundary = text.rfind(separator, minimum, hard_end)
                if boundary >= minimum:
                    end = boundary + len(separator)
                    break
        yield start, end, text[start:end]
        start = end


def _evidence_context_items(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for record in records:
        text = str(record.get("text") or "")
        if not text:
            continue
        source_id = record.get("source_id")
        metadata = {
            "source_id": source_id,
            "source_type": record.get("source_type"),
        }
        if record.get("source_type") == "website":
            metadata.update({
                "url": _display_url(record.get("final_url") or record.get("url")),
                "title": _compact_metadata(record.get("title"), 500),
                "language_hint": _compact_metadata(record.get("language_hint"), 40),
            })
        else:
            metadata.update({
                "file_name": _compact_metadata(record.get("file_name"), 260),
                "parsing_method": _compact_metadata(record.get("parsing_method"), 100),
            })
        chunks = list(_split_text_losslessly(text))
        for index, (start, end, chunk) in enumerate(chunks, start=1):
            items.append({
                **metadata,
                "chunk_id": f"{source_id}-c{index:04d}",
                "locator": f"characters {start + 1}-{end}",
                "text": chunk,
            })
    return items


def _queue_context_pages(
    state: OptimizedToolState,
    records: list[dict[str, Any]],
    candidates: list[dict[str, Any]],
) -> tuple[str | None, int, int, int]:
    evidence_items = _evidence_context_items(records)
    candidate_items = [
        _compact_candidate(candidate)
        for candidate in candidates
        if candidate.get("id") not in state.queued_candidate_ids
    ]
    pages: list[dict[str, Any]] = []
    current: dict[str, Any] = {"evidence": [], "candidate_links": []}

    def serialized_chars(page: dict[str, Any]) -> int:
        return len(json.dumps(page, ensure_ascii=False, separators=(",", ":")))

    def append_item(kind: str, item: dict[str, Any]) -> None:
        nonlocal current
        candidate_page = copy.deepcopy(current)
        candidate_page[kind].append(item)
        if (current["evidence"] or current["candidate_links"]) and serialized_chars(candidate_page) > MAX_CONTEXT_PAGE_CHARS:
            pages.append(current)
            current = {"evidence": [], "candidate_links": []}
            candidate_page = copy.deepcopy(current)
            candidate_page[kind].append(item)
        if serialized_chars(candidate_page) > MAX_CONTEXT_PAGE_CHARS:
            raise ValueError("context_item_exceeds_page_budget")
        current = candidate_page

    for item in evidence_items:
        append_item("evidence", item)
    for item in candidate_items:
        append_item("candidate_links", item)
    if current["evidence"] or current["candidate_links"]:
        pages.append(current)

    queued_chars = sum(serialized_chars(page) for page in pages)
    if state.context_chars_queued + queued_chars > MAX_TOTAL_CONTEXT_CHARS:
        raise ValueError("model_context_budget_exceeded")

    cursors = [f"ctx{state.next_context_page_index + index:04d}" for index in range(len(pages))]
    state.next_context_page_index += len(pages)
    for index, (cursor, page) in enumerate(zip(cursors, pages)):
        page["cursor"] = cursor
        page["next_cursor"] = cursors[index + 1] if index + 1 < len(cursors) else None
        state.context_pages[cursor] = page
        state.pending_context_cursors.append(cursor)
    state.queued_candidate_ids.update(
        candidate.get("id") for candidate in candidate_items if candidate.get("id")
    )
    state.context_chars_queued += queued_chars
    evidence_chars = sum(len(item["text"]) for item in evidence_items)
    _persist_context_delivery(state)
    return (cursors[0] if cursors else None), len(cursors), evidence_chars, len(candidate_items)


def _persist_context_delivery(state: OptimizedToolState) -> None:
    _write_json_atomic(state.workspace / "work" / "context-delivery.json", {
        "pages_queued": len(state.context_pages),
        "pages_delivered": len(state.delivered_context_cursors),
        "pages_remaining": len(state.pending_context_cursors),
        "context_chars_queued": state.context_chars_queued,
        "evidence_chars_delivered": state.evidence_chars_returned,
        "complete": not state.pending_context_cursors,
    })


def _candidate_score(entry: dict[str, Any]) -> int:
    searchable = " ".join(str(entry.get(name, "")) for name in ("url", "anchor_text", "surrounding_text")).lower()
    return sum(1 for keyword in CANDIDATE_KEYWORDS if keyword in searchable)


def _candidate_entries(workspace: Path) -> list[dict[str, Any]]:
    existing = _read_json(workspace / "work" / "link-candidates.json", {"links": []})
    by_url = {entry["url"]: entry for entry in existing.get("links", []) if entry.get("url")}
    next_index = max((int(e["id"][1:]) for e in by_url.values() if str(e.get("id", ""))[1:].isdigit()), default=0) + 1

    page_records = _page_records(workspace)
    fetched_urls = {
        value for page in page_records
        for value in (page.get("url"), page.get("final_url"), page.get("canonical_url"))
        if value
    }
    for page in page_records:
        for link in page.get("internal_links", []):
            if isinstance(link, str):
                url, text = link, ""
            else:
                url, text = link.get("href"), link.get("text", "")
            if not url or url in by_url or url in fetched_urls:
                continue
            by_url[url] = {
                "id": f"l{next_index:03d}",
                "url": url,
                "anchor_text": text,
                "surrounding_text": link.get("surrounding_text", "") if isinstance(link, dict) else "",
                "discovered_on": page["source_id"],
            }
            next_index += 1

    entries = sorted(by_url.values(), key=lambda e: e["id"])
    visible = sorted(entries, key=lambda e: (-_candidate_score(e), e["id"]))
    _write_json_atomic(workspace / "work" / "link-candidates.json", {
        "links": entries,
        "visible_count": len(visible),
        "omitted_count": 0,
    })
    return visible


def _source_summary(workspace: Path) -> dict[str, Any]:
    fetched = _read_json(workspace / "work" / "fetched-pages.json", {"pages": [], "failures": []})
    parsed = _read_json(workspace / "work" / "parsed-documents.json", {"documents": []})
    return {
        "pages_fetched": len(fetched.get("pages", [])),
        "page_failures": fetched.get("failures", []),
        "documents": [{
            "source_id": d.get("id"),
            "file_name": d.get("file_name"),
            "status": d.get("status"),
            "ocr_required": d.get("ocr_required", False),
            "reason": d.get("reason"),
        } for d in parsed.get("documents", [])],
    }


def build_optimized_tool_server(
    workspace: Path,
    package_dir: Path,
    telemetry: RunTelemetry,
):
    """Return (SDK MCP server, mutable run state, fully-qualified tool names)."""
    state = OptimizedToolState(workspace.resolve(), package_dir.resolve(), telemetry)
    state.bundle_limitations = []

    @tool("load_skill", "Load the complete Visit Finland SKILL.md and versioned extraction policy. Call this first.",
          EmptyToolInput.model_json_schema())
    @_idempotent_handler(state, "load_skill")
    async def load_skill(_args: dict[str, Any]):
        _parsed, invalid = _validated(EmptyToolInput, _args)
        if invalid:
            return invalid
        op = telemetry.start_operation("load_skill", "stage_0_bootstrap")
        try:
            skill_path = state.package_dir / "SKILL.md"
            skill_text = skill_path.read_text(encoding="utf-8")
            references = {}
            for name in [
                "general-curation.md", "accommodation-guidance.md", "field-semantics.md",
                "source-and-evidence-policy.md",
            ]:
                references[name] = (state.package_dir / "references" / name).read_text(encoding="utf-8")
            schemas = {
                "extraction": _read_json(state.package_dir / "schemas" / "accommodation-extraction.schema.json"),
                "schema_version": _read_json(state.package_dir / "schemas" / "schema-version.json"),
                "category_taxonomy": _read_json(state.package_dir / "schemas" / "datahub-categories.json"),
            }
            state.skill_loaded = True
            state.stage_reached = "stage_0_bootstrap"
            digest = hashlib.sha256(skill_text.encode("utf-8")).hexdigest()
            telemetry.finish_operation(op, skill_sha256=digest)
            telemetry.emit({"type": "skill_loaded", "skill_sha256": digest})
            return _json_result({
                "status": "ok", "workflow_mode": "coarse_tools", "skill_sha256": digest,
                "skill": skill_text, "references": references, "schemas": schemas,
            })
        except Exception as exc:
            telemetry.finish_operation(op, "error", error_code="skill_load_failed")
            state.forced_status, state.forced_error_code = "execution_failed", "skill_load_failed"
            return _json_result({"status": "error", "error_code": "skill_load_failed", "message": type(exc).__name__}, is_error=True)

    @tool("prepare_sources", "Validate the request, parse optional documents, fetch seed pages, and prepare lossless paginated context. Call after load_skill.",
          EmptyToolInput.model_json_schema())
    async def prepare_sources(_args: dict[str, Any]):
        _parsed, invalid = _validated(EmptyToolInput, _args)
        if invalid:
            return invalid
        if not state.skill_loaded:
            return _json_result({"status": "error", "error_code": "skill_not_loaded"}, is_error=True)
        if state.prepared:
            candidates = _candidate_entries(state.workspace)
            return _json_result({
                "status": "ok",
                "idempotent": True,
                "summary": _source_summary(state.workspace),
                "candidate_count": len(candidates),
                "context_complete": not state.pending_context_cursors,
                "next_context_cursor": state.pending_context_cursors[0] if state.pending_context_cursors else None,
                "message": "Sources are already prepared. Continue from the existing context cursor.",
                "next_action": {
                    "tool": "read_context_page" if state.pending_context_cursors else "fetch_selected_pages_or_record_scope",
                    "instruction": (
                        "Read every remaining context page in order. Do not call prepare_sources again."
                        if state.pending_context_cursors else
                        "All prepared context was delivered. Fetch relevant candidates or determine scope."
                    ),
                },
            })

        request_path = state.workspace / "input" / "request.json"
        manifest_path = state.workspace / "work" / "source-manifest.json"
        parsed_path = state.workspace / "work" / "parsed-documents.json"
        fetched_path = state.workspace / "work" / "fetched-pages.json"

        validation_op = telemetry.start_operation("validate_sources", "stage_1")
        validation = await _run_script(
            state, "validate_source_manifest_optimized.py",
            "--input", str(request_path),
            "--schema", str(state.package_dir / "schemas" / "analysis-request.schema.json"),
            "--output", str(manifest_path),
        )
        if validation["returncode"] != 0:
            telemetry.finish_operation(validation_op, "error", error_code="input_invalid")
            state.forced_status, state.forced_error_code = "input_invalid", (validation.get("error") or {}).get("error_code", "contract_validation_failed")
            state.stage_reached = "stage_1"
            return _json_result({"status": "input_invalid", "details": validation.get("error")}, is_error=True)
        telemetry.finish_operation(validation_op, **_metric_fields(validation.get("result")))
        state.stage_reached = "stage_1"

        manifest = _read_json(manifest_path, {})
        parse_op = telemetry.start_operation("parse_documents", "stage_2", documents_requested=len(manifest.get("documents", [])))
        if not manifest.get("documents"):
            _write_json_atomic(parsed_path, {"documents": []})
            parsing = {"returncode": 0, "result": {"status": "ok", "documents_parsed": 0, "empty_fast_path": True}}
        else:
            parsing = await _run_script(state, "parse_documents.py", "--manifest", str(manifest_path), "--output", str(parsed_path))
        telemetry.finish_operation(parse_op, "ok" if parsing["returncode"] == 0 else "error", **_metric_fields(parsing.get("result")))
        state.stage_reached = "stage_2"
        if parsing["returncode"] != 0:
            state.forced_status, state.forced_error_code = "execution_failed", "document_parser_error"
            return _json_result({"status": "error", "error_code": "document_parser_error"}, is_error=True)

        fetch_op = telemetry.start_operation("fetch_seed_pages", "stage_3", seeds=len(manifest.get("sources", [])))
        fetching = await _run_script(
            state, "fetch_pages_optimized.py",
            "--manifest", str(manifest_path),
            "--pages-dir", str(state.workspace / "work" / "pages"),
            "--output", str(fetched_path),
        )
        telemetry.finish_operation(fetch_op, "ok" if fetching["returncode"] == 0 else "error", **_metric_fields(fetching.get("result")))
        state.stage_reached = "stage_3"
        if fetching["returncode"] != 0:
            state.forced_status, state.forced_error_code = "execution_failed", "fetcher_error"
            return _json_result({"status": "error", "error_code": "fetcher_error"}, is_error=True)

        state.prepared = True
        records = [
            record for record in (_page_records(state.workspace) + _document_records(state.workspace))
            if str(record.get("text") or "")
        ]
        candidates = _candidate_entries(state.workspace)
        summary = _source_summary(state.workspace)
        if not records:
            state.forced_status = "no_usable_sources"
            state.forced_error_code = "documents_require_ocr" if any(d.get("ocr_required") for d in summary["documents"]) else "all_pages_inaccessible"
            return _json_result({
                "status": state.forced_status,
                "error_code": state.forced_error_code,
                "summary": summary,
                "terminal": True,
            })
        try:
            first_cursor, context_page_count, evidence_chars, queued_candidates = _queue_context_pages(state, records, candidates)
        except ValueError as exc:
            state.forced_status, state.forced_error_code = "execution_failed", str(exc)
            return _json_result({
                "status": "error", "error_code": str(exc),
                "message": "The complete evidence cannot fit safely in the configured model context budget.",
            }, is_error=True)
        fetch_op["evidence_chars_queued"] = evidence_chars
        return _json_result({
            "status": "ok",
            "summary": summary,
            "evidence_chars": evidence_chars,
            "candidate_count": queued_candidates,
            "context_page_count": context_page_count,
            "next_context_cursor": first_cursor,
            "context_complete": context_page_count == 0,
            "limitations": [],
            "capabilities": {
                "static_html": True, "javascript_rendering": False, "general_web_search": False,
                "documents": ["pdf", "docx"], "retrieval_scope": "caller-approved domains only",
            },
            "next_action": {
                "tool": "read_context_page",
                "cursor": first_cursor,
                "instruction": "Read every context page in order until context_complete is true. Then fetch relevant candidate IDs or record scope.",
            },
        })

    @tool("read_context_page", "Read the next exact, lossless page of prepared website/document evidence and candidate links. Read every page in cursor order before fetching more pages or recording scope.",
          ReadContextPageInput.model_json_schema())
    @_idempotent_handler(state, "read_context_page")
    async def read_context_page(args: dict[str, Any]):
        parsed, invalid = _validated(ReadContextPageInput, args)
        if invalid:
            return invalid
        if state.forced_status is not None:
            return _json_result({
                "status": state.forced_status,
                "error_code": state.forced_error_code,
                "terminal": True,
            })
        if not state.prepared:
            return _json_result({"status": "error", "error_code": "sources_not_prepared"}, is_error=True)
        cursor = parsed["cursor"]
        if cursor not in state.context_pages:
            return _json_result({"status": "error", "error_code": "unknown_context_cursor"}, is_error=True)
        expected = state.pending_context_cursors[0] if state.pending_context_cursors else None
        if cursor != expected:
            return _json_result({
                "status": "error",
                "error_code": "context_cursor_out_of_order",
                "expected_cursor": expected,
            }, is_error=True)

        page = state.context_pages[cursor]
        evidence_chars = sum(len(str(item.get("text", ""))) for item in page["evidence"])
        context_complete = len(state.pending_context_cursors) == 1
        payload = {
            "status": "ok",
            "cursor": cursor,
            "evidence": page["evidence"],
            "candidate_links": page["candidate_links"],
            "next_cursor": page["next_cursor"],
            "context_complete": context_complete,
            "remaining_context_pages": len(state.pending_context_cursors) - 1,
            "next_action": {
                "tool": "read_context_page" if not context_complete else "fetch_selected_pages_or_record_scope",
                "cursor": page["next_cursor"],
                "instruction": (
                    "Read the next context page before any scope or extraction decision."
                    if not context_complete else
                    "All currently prepared evidence and candidates have been delivered. Fetch relevant candidate IDs if needed; otherwise record scope."
                ),
            },
        }
        response = _json_result(payload)
        if len(response["content"][0]["text"]) > MAX_MCP_RESPONSE_CHARS:
            state.forced_status, state.forced_error_code = "execution_failed", "pagination_response_too_large"
            return _json_result({
                "status": "error",
                "error_code": "pagination_response_too_large",
            }, is_error=True)

        op = telemetry.start_operation("read_context_page", "stage_3", cursor=cursor)
        state.pending_context_cursors.pop(0)
        state.delivered_context_cursors.add(cursor)
        state.evidence_chars_returned += evidence_chars
        _persist_context_delivery(state)
        telemetry.finish_operation(
            op,
            evidence_chunks=len(page["evidence"]),
            evidence_chars_returned=evidence_chars,
            candidate_links=len(page["candidate_links"]),
            context_complete=context_complete,
        )
        return response

    @tool("fetch_selected_pages", "Fetch relevant same-site pages by candidate link ID. Use semantic judgment; never invent IDs. At most two bounded calls are allowed.",
          FetchSelectedPagesInput.model_json_schema())
    @_idempotent_handler(state, "fetch_selected_pages")
    async def fetch_selected_pages(args: dict[str, Any]):
        parsed, invalid = _validated(FetchSelectedPagesInput, args)
        if invalid:
            return invalid
        args = parsed
        if state.forced_status is not None:
            return _json_result({
                "status": state.forced_status,
                "error_code": state.forced_error_code,
                "terminal": True,
            })
        if not state.prepared:
            return _json_result({"status": "error", "error_code": "sources_not_prepared"}, is_error=True)
        if state.pending_context_cursors:
            return _json_result({
                "status": "error",
                "error_code": "context_not_fully_delivered",
                "next_context_cursor": state.pending_context_cursors[0],
            }, is_error=True)
        if state.expansion_calls >= MAX_EXPANSION_CALLS:
            return _json_result({"status": "error", "error_code": "expansion_call_limit_reached"}, is_error=True)
        requested_ids = list(dict.fromkeys(args.get("link_ids", [])))
        new_link_ids = set(requested_ids) - state.selected_link_ids
        current_page_count = len(_page_records(state.workspace))
        call_limit = MAX_SELECTED_PER_CALL if state.expansion_calls == 0 else MAX_RECOVERY_SELECTED
        if (
            len(requested_ids) > call_limit
            or len(state.selected_link_ids | set(requested_ids)) > MAX_SELECTED_PER_RUN
            or current_page_count + len(new_link_ids) > MAX_TOTAL_PAGES
        ):
            return _json_result({"status": "error", "error_code": "page_budget_exceeded"}, is_error=True)
        candidates_doc = _read_json(state.workspace / "work" / "link-candidates.json", {"links": []})
        candidates = {entry["id"]: entry for entry in candidates_doc.get("links", [])}
        unknown = [link_id for link_id in requested_ids if link_id not in candidates]
        if unknown:
            return _json_result({"status": "error", "error_code": "unknown_link_ids", "link_ids": unknown}, is_error=True)

        before_ids = {p["source_id"] for p in _page_records(state.workspace)}
        op = telemetry.start_operation("fetch_selected_pages", "stage_3", requested=len(requested_ids), recovery_pass=state.expansion_calls > 0)
        command_args = [
            "--manifest", str(state.workspace / "work" / "source-manifest.json"),
            "--pages-dir", str(state.workspace / "work" / "pages"),
            "--output", str(state.workspace / "work" / "fetched-pages.json"),
        ]
        for link_id in requested_ids:
            command_args.extend(["--url", candidates[link_id]["url"]])
        fetching = await _run_script(state, "fetch_pages_optimized.py", *command_args)
        state.expansion_calls += 1
        state.selected_link_ids.update(requested_ids)
        telemetry.finish_operation(op, "ok" if fetching["returncode"] == 0 else "error", **_metric_fields(fetching.get("result")))
        if fetching["returncode"] != 0:
            state.forced_status, state.forced_error_code = "execution_failed", "fetcher_error"
            return _json_result({"status": "error", "error_code": "fetcher_error"}, is_error=True)

        after_records = _page_records(state.workspace)
        new_ids = {p["source_id"] for p in after_records} - before_ids
        records = [p for p in after_records if p["source_id"] in new_ids]
        next_candidates = _candidate_entries(state.workspace)
        unseen_candidates = [c for c in next_candidates if c["id"] not in state.selected_link_ids]
        try:
            first_cursor, context_page_count, evidence_chars, queued_candidates = _queue_context_pages(
                state, records, unseen_candidates,
            )
        except ValueError as exc:
            state.forced_status, state.forced_error_code = "execution_failed", str(exc)
            return _json_result({
                "status": "error", "error_code": str(exc),
                "message": "The complete evidence cannot fit safely in the configured model context budget.",
            }, is_error=True)
        op["evidence_chars_queued"] = evidence_chars
        return _json_result({
            "status": "ok", "fetch_result": fetching.get("result"),
            "new_evidence_chars": evidence_chars,
            "additional_candidate_count": queued_candidates,
            "context_page_count": context_page_count,
            "next_context_cursor": first_cursor,
            "context_complete": context_page_count == 0,
            "remaining": {
                "expansion_calls": MAX_EXPANSION_CALLS - state.expansion_calls,
                "selected_pages": MAX_SELECTED_PER_RUN - len(state.selected_link_ids),
                "next_call_pages": MAX_RECOVERY_SELECTED if state.expansion_calls == 1 else 0,
            },
            "limitations": [],
            "next_action": {
                "tool": "read_context_page" if first_cursor else "record_scope",
                "cursor": first_cursor,
                "instruction": (
                    "Read every newly prepared context page in order before determining scope. "
                    "Use a recovery fetch only after all pages are read and only if they reveal a specific essential missing page."
                    if first_cursor else
                    "No new context was produced. Determine and record scope from the context already delivered."
                ),
            },
        })

    @tool("record_scope", "Store a scope decision. Status must be resolved, or scope_ambiguous when multiple products or an unclear boundary prevent choosing one.",
          ScopeInput.model_json_schema())
    @_idempotent_handler(state, "record_scope")
    async def record_scope(args: dict[str, Any]):
        if state.forced_status is not None:
            return _json_result({
                "status": state.forced_status,
                "error_code": state.forced_error_code,
                "terminal": True,
            })
        parsed, invalid = _validated(ScopeInput, args)
        if invalid:
            return invalid
        args = parsed
        if not state.prepared:
            return _json_result({"status": "error", "error_code": "sources_not_prepared"}, is_error=True)
        if state.pending_context_cursors:
            return _json_result({
                "status": "error",
                "error_code": "context_not_fully_delivered",
                "next_context_cursor": state.pending_context_cursors[0],
            }, is_error=True)
        _mark_scope_reasoning_started(state)
        decision = args.get("scope_decision")
        if not isinstance(decision, dict) or decision.get("status") not in {"resolved", "scope_ambiguous"}:
            return _json_result({"status": "error", "error_code": "invalid_scope_status"}, is_error=True)
        required_lists = ("excluded", "additional_products", "out_of_type_facilities")
        if any(not isinstance(decision.get(name), list) for name in required_lists):
            return _json_result({"status": "error", "error_code": "invalid_scope_shape", "required_lists": list(required_lists)}, is_error=True)
        known_sources = {r["source_id"] for r in _page_records(state.workspace)} | {r["source_id"] for r in _document_records(state.workspace)}
        if decision["status"] == "resolved":
            product = decision.get("product")
            pages = product.get("pages") if isinstance(product, dict) else None
            if not isinstance(pages, list) or not product.get("summary"):
                return _json_result({"status": "error", "error_code": "resolved_scope_requires_product"}, is_error=True)
            unknown_sources = [source_id for source_id in pages if source_id not in known_sources]
            if unknown_sources:
                return _json_result({"status": "error", "error_code": "unknown_scope_sources", "source_ids": unknown_sources}, is_error=True)
            decision["error_code"] = None
        else:
            if decision.get("error_code") not in {"multiple_accommodation_products", "product_boundary_unclear"}:
                return _json_result({"status": "error", "error_code": "invalid_scope_error_code"}, is_error=True)

        op = telemetry.start_operation("record_scope", "stage_4")
        _write_json_atomic(state.workspace / "work" / "scope-decision.json", decision)
        state.scope_status = decision["status"]
        state.stage_reached = "stage_4"
        telemetry.finish_operation(op, scope_status=state.scope_status)
        review_required = state.scope_status == "scope_ambiguous"
        return _json_result({
            "status": "ok",
            "scope_status": state.scope_status,
            "continue_to_extraction": True,
            "review_required": review_required,
            "instruction": (
                "Continue extraction without merging candidate products. Mark fields affected by "
                "scope uncertainty as review or missing and explain alternatives."
                if review_required else "Continue to extraction."
            ),
        })

    @tool("submit_extraction", "Store and deterministically validate a complete schema-guided extraction. One repaired resubmission is allowed.",
          ExtractionInput.model_json_schema())
    @_idempotent_handler(state, "submit_extraction")
    async def submit_extraction(args: dict[str, Any]):
        parsed, invalid = _validated(ExtractionInput, args)
        if invalid:
            return invalid
        args = parsed
        if state.pending_context_cursors:
            return _json_result({
                "status": "error",
                "error_code": "context_not_fully_delivered",
                "next_context_cursor": state.pending_context_cursors[0],
            }, is_error=True)
        if state.scope_status not in {"resolved", "scope_ambiguous"}:
            return _json_result({"status": "error", "error_code": "scope_not_resolved"}, is_error=True)
        if state.extraction_submissions >= 2:
            return _json_result({"status": "error", "error_code": "repair_limit_reached"}, is_error=True)
        draft = args.get("extraction")
        if not isinstance(draft, dict):
            return _json_result({"status": "error", "error_code": "invalid_extraction_shape"}, is_error=True)

        draft_path = state.workspace / "work" / "extraction-draft.json"
        result_path = state.workspace / "work" / "staging-output" / "result.json"
        stage = "stage_5" if state.extraction_submissions == 0 else "stage_6"
        op = telemetry.start_operation("submit_extraction", stage, submission=state.extraction_submissions + 1)
        _write_json_atomic(draft_path, draft)
        state.stage_reached = "stage_5"
        validation = await _run_script(
            state, "validate_extraction.py",
            "--input", str(draft_path),
            "--schema", str(state.package_dir / "schemas" / "accommodation-extraction.schema.json"),
            "--output", str(result_path),
        )
        state.extraction_submissions += 1
        state.stage_reached = "stage_6"
        result = _read_json(result_path, {"validation": {"valid": False, "errors": [{"rule_id": "validator_no_result"}], "warnings": []}})
        diagnostics = result.get("validation", {})
        valid = diagnostics.get("valid") is True
        telemetry.finish_operation(op, "ok" if valid else "invalid", valid=valid, error_count=len(diagnostics.get("errors", [])), warning_count=len(diagnostics.get("warnings", [])))

        return _json_result({
            "status": "valid" if valid else "invalid",
            "validation": diagnostics,
            "repair_allowed": not valid and state.extraction_submissions < 2,
            "submission": state.extraction_submissions,
        })

    @tool("finalize_outputs", "Build and verify canonical JSON and Excel, and store Claude's review-report prose. Call only after valid extraction.",
          FinalizeInput.model_json_schema())
    @_idempotent_handler(state, "finalize_outputs")
    async def finalize_outputs(args: dict[str, Any]):
        parsed, invalid = _validated(FinalizeInput, args)
        if invalid:
            return invalid
        args = parsed
        result_path = state.workspace / "work" / "staging-output" / "result.json"
        result = _read_json(result_path, {})
        if result.get("validation", {}).get("valid") is not True:
            return _json_result({"status": "error", "error_code": "extraction_not_valid"}, is_error=True)
        if state.finalized:
            artifacts = ["result.json", "canonical-product.json", "result.xlsx", "review-report.md"]
            if (state.workspace / "work" / "staging-output" / "scope-decision.json").is_file():
                artifacts.append("scope-decision.json")
            return _json_result({"status": "ok", "idempotent": True, "artifacts": artifacts})

        staging = state.workspace / "work" / "staging-output"
        canonical_path = staging / "canonical-product.json"
        xlsx_path = staging / "result.xlsx"
        report_path = staging / "review-report.md"
        scope_path = state.workspace / "work" / "scope-decision.json"
        staged_scope_path = staging / "scope-decision.json"
        op = telemetry.start_operation("finalize_outputs", "stage_7")
        canonical = await _run_script(state, "build_canonical_record.py", "--input", str(result_path), "--output", str(canonical_path))
        if canonical["returncode"] != 0:
            telemetry.finish_operation(op, "error", error_code="canonical_export_error")
            state.forced_status, state.forced_error_code = "execution_failed", "canonical_export_error"
            return _json_result({"status": "error", "error_code": "canonical_export_error"}, is_error=True)
        workbook = await _run_script(state, "export_xlsx.py", "--input", str(result_path), "--output", str(xlsx_path))
        if workbook["returncode"] != 0:
            telemetry.finish_operation(op, "error", error_code="export_error")
            state.forced_status, state.forced_error_code = "execution_failed", "export_error"
            return _json_result({"status": "error", "error_code": "export_error"}, is_error=True)

        report = args["review_report"].strip()
        if not report.startswith("#"):
            report = "# Visit Finland DataHub preparation review\n\n" + report
        report_path.write_text(
            with_deterministic_appendix(
                state.workspace, report, result.get("validation", {}), state.bundle_limitations,
            ),
            encoding="utf-8",
        )
        scope_decision = _read_json(scope_path, {})
        if scope_decision.get("status") == "scope_ambiguous":
            _write_json_atomic(staged_scope_path, scope_decision)

        try:
            result_document = _read_json(result_path)
            canonical_document = _read_json(canonical_path)
            if not isinstance(result_document, dict) or not isinstance(result_document.get("fields"), dict):
                raise ValueError("result.json has no fields object")
            if not isinstance(canonical_document, dict) or not isinstance(canonical_document.get("fields"), dict):
                raise ValueError("canonical-product.json has no fields object")
            extraction_schema = _read_json(
                state.package_dir / "schemas" / "accommodation-extraction.schema.json",
            )
            canonical_schema = _read_json(
                state.package_dir / "schemas" / "canonical-accommodation.schema.json",
            )
            jsonschema.Draft202012Validator(
                extraction_schema, format_checker=jsonschema.FormatChecker(),
            ).validate(result_document)
            jsonschema.Draft202012Validator(
                canonical_schema, format_checker=jsonschema.FormatChecker(),
            ).validate(canonical_document)
            book = load_workbook(xlsx_path, read_only=True)
            try:
                if not book.sheetnames:
                    raise ValueError("workbook has no sheets")
            finally:
                book.close()
            if report_path.stat().st_size < 20:
                raise ValueError("review report is empty")
            artifact_paths = [result_path, canonical_path, xlsx_path, report_path]
            if staged_scope_path.is_file():
                artifact_paths.append(staged_scope_path)
            state.artifact_hashes = {
                path.name: hashlib.sha256(path.read_bytes()).hexdigest()
                for path in artifact_paths
            }
        except Exception as exc:
            telemetry.finish_operation(op, "error", error_code="output_verification_failed")
            state.forced_status, state.forced_error_code = "execution_failed", "output_verification_failed"
            return _json_result({"status": "error", "error_code": "output_verification_failed", "message": type(exc).__name__}, is_error=True)

        state.finalized = True
        state.stage_reached = "stage_7"
        artifacts = ["result.json", "canonical-product.json", "result.xlsx", "review-report.md"]
        if staged_scope_path.is_file():
            artifacts.append("scope-decision.json")
        telemetry.finish_operation(op, artifacts=len(artifacts))
        return _json_result({
            "status": "ok", "artifacts": artifacts,
            "artifact_hashes": state.artifact_hashes,
            "message": "All staged artifacts verified. End the run now.",
        })

    server = create_sdk_mcp_server(
        name=SERVER_NAME,
        version="1.0.0",
        tools=[load_skill, prepare_sources, read_context_page, fetch_selected_pages, record_scope, submit_extraction, finalize_outputs],
    )
    state.handlers = {
        "load_skill": load_skill.handler,
        "prepare_sources": prepare_sources.handler,
        "read_context_page": read_context_page.handler,
        "fetch_selected_pages": fetch_selected_pages.handler,
        "record_scope": record_scope.handler,
        "submit_extraction": submit_extraction.handler,
        "finalize_outputs": finalize_outputs.handler,
    }
    return server, state, list(ALLOWED_TOOL_NAMES)







