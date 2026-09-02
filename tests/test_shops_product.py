import json
import subprocess
import sys

import jsonschema
import openpyxl
import pytest

import app.agent_tools as agent_tools
from app.telemetry import RunTelemetry
from conftest import PACKAGE_DIR, SCHEMAS_DIR, SCRIPTS_DIR


def envelope(status, value=None, evidence=None):
    result = {"status": status}
    if value is not None:
        result["value"] = value
    if evidence is not None:
        result["evidence"] = evidence
    return result


def shops_draft(category_id="artisan"):
    return {
        "run_metadata": {
            "schema_name": "poc-shops-schema-v1",
            "schema_version": "0.2.0-poc",
            "category_taxonomy_version": "poc-subset-0.3",
            "generated_at": "2026-09-02T00:00:00Z",
        },
        "product_type": "shops",
        "fields": {
            "name": {"fi": envelope("missing"), "en": envelope("missing")},
            "description": {"fi": envelope("missing"), "en": envelope("missing")},
            "address": envelope("missing"),
            "coordinates": envelope("missing"),
            "contact": envelope("missing"),
            "categories": envelope("review", [category_id]),
            "accessibility": envelope("missing"),
            "sustainability_label": envelope("missing"),
            "stf_status": envelope("missing"),
            "opening_hours": envelope("missing"),
            "languages_spoken": envelope("missing"),
            "images": envelope("not_in_scope"),
        },
    }


def run_shops_validator(workspace, draft):
    input_path = workspace / "work" / "extraction-draft.json"
    output_path = workspace / "work" / "staging-output" / "result.json"
    input_path.write_text(json.dumps(draft), encoding="utf-8")
    result = subprocess.run(
        [
            sys.executable,
            str(SCRIPTS_DIR / "validate_extraction.py"),
            "--input",
            str(input_path),
            "--schema",
            str(SCHEMAS_DIR / "shops-extraction.schema.json"),
            "--output",
            str(output_path),
        ],
        capture_output=True,
        text=True,
    )
    return result, json.loads(output_path.read_text(encoding="utf-8"))


def test_taxonomy_uses_plain_ids_with_explicit_groups():
    taxonomy = json.loads((SCHEMAS_DIR / "datahub-categories.json").read_text(encoding="utf-8"))
    assert taxonomy["taxonomy_version"] == "poc-subset-0.3"
    assert all("." not in category["id"] for category in taxonomy["categories"])
    assert {category["group"] for category in taxonomy["categories"]} == {
        "accommodation", "amenity", "shops",
    }


def test_request_contract_accepts_shops_and_rejects_other_product_types():
    schema = json.loads((SCHEMAS_DIR / "analysis-request.schema.json").read_text(encoding="utf-8"))
    request = {
        "contract_version": "v1",
        "product_type": "shops",
        "website_urls": ["https://example.fi/shop"],
    }
    jsonschema.Draft202012Validator(schema).validate(request)

    request["product_type"] = "restaurants"
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.Draft202012Validator(schema).validate(request)


def test_shops_schema_contains_only_confirmed_shared_fields():
    schema = json.loads((SCHEMAS_DIR / "shops-extraction.schema.json").read_text(encoding="utf-8"))
    fields = schema["properties"]["fields"]
    assert set(fields["required"]) == {
        "name",
        "description",
        "address",
        "coordinates",
        "contact",
        "categories",
        "accessibility",
        "sustainability_label",
        "stf_status",
        "opening_hours",
        "languages_spoken",
        "images",
    }
    assert {"capacity", "pricing", "booking_url", "availability", "amenities"}.isdisjoint(
        fields["properties"]
    )


def test_shops_validator_accepts_shop_category_and_metadata(workspace):
    result, output = run_shops_validator(workspace, shops_draft())
    assert result.returncode == 0, result.stderr
    assert output["product_type"] == "shops"
    assert output["run_metadata"]["schema_name"] == "poc-shops-schema-v1"
    assert output["run_metadata"]["category_taxonomy_version"] == "poc-subset-0.3"
    assert not any(w["rule_id"] == "unknown_category" for w in output["validation"]["warnings"])


def test_shops_validator_normalizes_legacy_prefixed_category(workspace):
    result, output = run_shops_validator(workspace, shops_draft("shops.shopping_center"))
    assert result.returncode == 0
    assert output["fields"]["categories"]["value"] == ["shopping_center"]
    assert not any(w["rule_id"] == "unknown_category" for w in output["validation"]["warnings"])


def test_shops_validator_flags_cross_product_category(workspace):
    result, output = run_shops_validator(workspace, shops_draft("hotel"))
    assert result.returncode == 0
    warning = next(w for w in output["validation"]["warnings"] if w["rule_id"] == "unknown_category")
    assert warning["values"] == ["hotel"]


def test_shops_export_uses_shops_sheet_and_shared_columns_only(workspace):
    result_path = workspace / "work" / "staging-output" / "result.json"
    result_path.parent.mkdir(parents=True, exist_ok=True)
    result_path.write_text(json.dumps(shops_draft()), encoding="utf-8")
    output_path = result_path.with_name("result.xlsx")

    completed = subprocess.run(
        [
            sys.executable,
            str(SCRIPTS_DIR / "export_xlsx.py"),
            "--input",
            str(result_path),
            "--output",
            str(output_path),
        ],
        capture_output=True,
        text=True,
    )
    assert completed.returncode == 0, completed.stderr

    workbook = openpyxl.load_workbook(output_path)
    assert workbook.sheetnames == ["Shops"]
    headers = [cell.value for cell in next(workbook["Shops"].iter_rows(min_row=1, max_row=1))]
    assert "opening_hours" in headers
    assert "capacity" not in headers
    assert "booking_url" not in headers


@pytest.mark.asyncio
async def test_load_skill_returns_shops_schema_and_guidance(workspace):
    request_path = workspace / "input" / "request.json"
    request_path.write_text(json.dumps({
        "contract_version": "v1",
        "product_type": "shops",
        "website_urls": ["https://example.fi/shop"],
        "document_files": [],
    }), encoding="utf-8")

    _server, state, _allowed = agent_tools.build_optimized_tool_server(
        workspace, PACKAGE_DIR, RunTelemetry("test"),
    )
    response = await state.handlers["load_skill"]({})
    payload = json.loads(response["content"][0]["text"])

    assert payload["product_type"] == "shops"
    assert payload["schemas"]["extraction"]["$id"] == "shops-extraction.schema.json"
    assert payload["schemas"]["schema_version"]["schema_name"] == "poc-shops-schema-v1"
    assert "shops-guidance.md" in payload["references"]
    assert "accommodation-guidance.md" not in payload["references"]


def test_shop_candidate_ranking_prefers_shop_links():
    generic = {"url": "https://example.fi/company", "anchor_text": "Company", "surrounding_text": ""}
    relevant = {
        "url": "https://example.fi/boutique",
        "anchor_text": "Local products shop",
        "surrounding_text": "Opening hours and store location",
    }
    assert agent_tools._candidate_score(relevant, "shops") > agent_tools._candidate_score(generic, "shops")
